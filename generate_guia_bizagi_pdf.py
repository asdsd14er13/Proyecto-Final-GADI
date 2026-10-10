import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_apa_builder import (
    create_document, add_p_apa, add_h1, add_h2, add_h3,
    add_table_apa, set_cell_margins, set_cell_borders, set_cell_shading
)

def add_callout_box(doc, title, text, bg_color="FEE2E2", border_color="EF4444", text_color="991B1B"):
    """Inserta una caja destacada (tipo callout) con borde lateral y sombreado."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    set_cell_shading(cell, bg_color)
    
    border_left = {'val': 'single', 'sz': 24, 'color': border_color}
    set_cell_borders(cell, top=None, bottom=None, left=border_left, right=None)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    r_tit = p.add_run(title + "\n")
    r_tit.bold = True
    r_tit.font.name = 'Times New Roman'
    r_tit.font.size = Pt(10.5)
    r_tit.font.color.rgb = RGBColor.from_string(text_color)
    
    r_txt = p.add_run(text)
    r_txt.font.name = 'Times New Roman'
    r_txt.font.size = Pt(10)
    r_txt.font.color.rgb = RGBColor(30, 41, 59)
    
    # Párrafo en blanco después
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def build_guia_docx():
    doc = create_document()
    
    # PORTADA / ENCABEZADO FORMAL
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(12)
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("UNIVERSIDAD NACIONAL FEDERICO VILLARREAL\nFACULTAD DE INGENIERÍA ELECTRÓNICA E INFORMÁTICA")
    r_inst.bold = True
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(12)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("ESCUELA PROFESIONAL DE INGENIERÍA INFORMÁTICA\nGESTIÓN Y ADMINISTRACIÓN DE DATOS E INFORMACIÓN (GADI)")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    
    add_h1(doc, "MANUAL DE ENCARGO TÉCNICO:\nMODELADO DE PROCESOS AS-IS Y TO-BE EN BIZAGI")
    
    # Metadata del encargo
    meta_headers = ["Campo", "Detalle de Especificación"]
    meta_rows = [
        ["Proyecto", "Sistema Inteligente de Evaluación Adaptativa y Personalización del Aprendizaje Universitario"],
        ["Destinatario", "Miembro del Equipo de Desarrollo / Modelador de Procesos BPMN"],
        ["Estándar", "BPMN 2.0 (Business Process Model and Notation) en Bizagi Process Modeler"],
        ["Entregables Requeridos", "1. Archivo nativo Bizagi (.bpm) con 2 diagramas.\n2. Dos exportaciones en alta resolución (.png a 300 DPI) para informe final APA 7."]
    ]
    add_table_apa(doc, "1", "Ficha Técnica del Encargo de Modelado", meta_headers, meta_rows, [2.0, 4.5])
    
    add_h2(doc, "1. Propósito y Lineamientos Generales")
    add_p_apa(doc, 
        "El presente manual formaliza las directrices para el modelado de los diagramas de procesos de negocio correspondientes al Capítulo I del proyecto de fin de curso. El modelador deberá traducir con total rigurosidad la secuencia cronológica de 16 semanas, los roles institucionales y los hitos de control descritos en el informe, asegurando consistencia entre la narrativa teórica y la representación visual en Bizagi.",
        bold_prefix="Objetivo: "
    )
    add_p_apa(doc, 
        "En Bizagi Process Modeler, cree un único archivo de proyecto denominado MODELADO_PROCESOS_GADI_BIZAGI.bpm. Dicho archivo deberá contener dos pestañas de diagramas claramente diferenciadas: la primera nombrada '1. Proceso Actual (AS-IS)' y la segunda nombrada '2. Proceso Futuro (TO-BE)'. Todos los flujos de secuencia deben trazarse ortogonalmente (ángulos de 90°), evitando cruces de líneas y asignando a cada tarea su tipo formal (Usuario, Manual o Servicio).",
        bold_prefix="Configuración Inicial: "
    )
    
    # SECCIÓN AS-IS
    add_h2(doc, "2. Especificación del Diagrama 1: Proceso Actual (AS-IS)")
    add_p_apa(doc,
        "El diagrama AS-IS refleja el estado operativo actual de la FIEI-UNFV: un flujo fragmentado, manual y reactivo donde la facultad permanece sin visibilidad del rendimiento durante la primera mitad del ciclo lectivo y las tutorías se ejecutan de manera tardía e indiferenciada.",
        bold_prefix="Contexto Operativo: "
    )
    
    add_p_apa(doc,
        "El diagrama debe contener un único Pool titulado 'PROCESO DE SEGUIMIENTO ACADÉMICO Y RETENCIÓN ESTUDIANTIL (AS-IS)', dividido estrictamente en 3 carriles (Lanes) horizontales: (1) Estudiante FIEI, (2) Docente de Asignatura y (3) Comité de Tutoría.",
        bold_prefix="Estructura de Contenedores: "
    )
    
    # Tabla AS-IS
    as_is_headers = ["N°", "Carril (Lane)", "Tipo BPMN", "Texto Exacto de la Tarea / Evento", "Salida / Artefacto"]
    as_is_rows = [
        ["0", "Estudiante FIEI", "Inicio (Círculo verde)", "Inicio del Semestre Académico", "Apertura de ciclo"],
        ["1", "Estudiante FIEI", "Tarea de Usuario", "1. Realizar matrícula regular (Semana 0)", "Ficha de matrícula básica"],
        ["2", "Estudiante FIEI", "Tarea de Usuario", "2. Asistir a clases y rendir evaluaciones continuas (Sem. 1-7)", "Prácticas rendidas"],
        ["3", "Docente Asignatura", "Tarea Manual", "3. Registrar calificaciones en hojas locales (Sem. 1-7)", "Excel personal en laptop"],
        ["4", "Estudiante FIEI", "Tarea de Usuario", "4. Rendir exámenes parciales (Semana 8)", "Pruebas físicas"],
        ["5", "Docente Asignatura", "Tarea de Usuario", "5. Calificar pruebas y digitar notas en intranet (Sem. 9-10)", "Acta parcial en intranet"],
        ["6", "Comité de Tutoría", "Tarea de Usuario", "6. Revisar actas y filtrar promedios < 10.5 manualmente (Sem. 10-11)", "Lista física de observados"],
        ["G1", "Comité de Tutoría", "Compuerta XOR", "¿Estudiante con promedio < 10.5?", "Bifurcación SÍ / NO"],
        ["7", "Comité de Tutoría", "Tarea Manual", "7. Emitir y enviar citación genérica en papel/correo (Sem. 11-12)", "Citación estándar"],
        ["8", "Estudiante FIEI", "Tarea de Usuario", "8. Asistir a sesión de tutoría tardía (Semana 12)", "Firma de asistencia"],
        ["9", "Estudiante FIEI", "Tarea de Usuario", "9. Rendir exámenes finales (Semana 16)", "Pruebas finales"],
        ["10", "Docente Asignatura", "Tarea de Usuario", "10. Digitar y cerrar actas finales del ciclo (Semana 16)", "Acta final cerrada"],
        ["Fin", "Docente Asignatura", "Fin (Círculo rojo)", "Cierre de ciclo con reprobación, rezago y pérdida de gratuidad", "Consumación de fracaso"]
    ]
    add_table_apa(doc, "2", "Inventario Detallado de Elementos del Proceso Actual (AS-IS)", as_is_headers, as_is_rows, [0.4, 1.3, 1.3, 2.2, 1.3])
    
    add_h3(doc, "Anotaciones Obligatorias de Cuellos de Botella (Dolores de Gestión)")
    add_p_apa(doc, "El modelador debe vincular obligatoriamente las siguientes cuatro anotaciones BPMN en recuadros de color rojo claro (#FEE2E2) con línea punteada hacia sus tareas respectivas:", first_indent=False)
    
    add_callout_box(doc, "⚡ DOLOR 1: Silos de Información (Conectar a Tarea 3)", 
                    "Las notas de evaluaciones continuas se almacenan en hojas de cálculo aisladas en computadoras docentes. La facultad opera a oscuras sin visibilidad institucional durante las primeras 7 semanas de clases.",
                    bg_color="FEE2E2", border_color="EF4444", text_color="991B1B")
    
    add_callout_box(doc, "⚡ DOLOR 2: Latencia Informacional Crítica (Conectar a Tarea 5)", 
                    "Retraso estructural de 9 a 10 semanas desde el inicio del semestre para que las autoridades y tutores dispongan de calificaciones consolidadas en el sistema oficial.",
                    bg_color="FEE2E2", border_color="EF4444", text_color="991B1B")
                    
    add_callout_box(doc, "⚡ DOLOR 3: Detección Reactiva 'Post-Mortem' (Conectar a Tarea 6)", 
                    "La identificación de estudiantes en riesgo se efectúa manualmente revisando promedios aritméticos simples (< 10.5) cuando el estudiante ya reprobó el examen parcial sin capacidad de reacción.",
                    bg_color="FEE2E2", border_color="EF4444", text_color="991B1B")
                    
    add_callout_box(doc, "⚡ DOLOR 4: Intervención Uniforme e Ineficaz (Conectar a Tarea 7)", 
                    "La citación a tutoría se emite tardíamente en semanas 11 y 12 de forma genérica, sin identificar si el rezago proviene de sobrecarga laboral, dificultades socioeconómicas o deficiencia matemática.",
                    bg_color="FEE2E2", border_color="EF4444", text_color="991B1B")

    # SECCIÓN TO-BE
    add_h2(doc, "3. Especificación del Diagrama 2: Proceso Futuro Optimizado (TO-BE)")
    add_p_apa(doc,
        "El diagrama TO-BE formaliza la reingeniería del proceso mediante la inclusión de la Plataforma de Business Intelligence y Machine Learning. El flujo pasa de ser manual y reactivo a un ciclo proactivo, continuo y adaptativo basado en dos momentos de evaluación (M1 y M2).",
        bold_prefix="Contexto de Reingeniería: "
    )
    
    add_p_apa(doc,
        "El diagrama debe contener un único Pool titulado 'PROCESO INTELIGENTE DE SEGUIMIENTO ADAPTATIVO Y RETENCIÓN ESTUDIANTIL CON BI + IA (TO-BE)', organizado en 4 carriles (Lanes) horizontales: (1) Estudiante FIEI, (2) Docente de Asignatura / Docente Tutor, (3) Dirección de Orientación y Tutoría (Comité / Directores) y (4) Plataforma Analítica Automatizada (BI + IA).",
        bold_prefix="Estructura de Contenedores: "
    )
    
    to_be_headers = ["N°", "Carril (Lane)", "Tipo BPMN", "Texto Exacto de la Tarea / Evento", "Artefacto / Tecnología"]
    to_be_rows = [
        ["0", "Estudiante FIEI", "Inicio (Círculo verde)", "Apertura de Proceso de Matrícula", "Convocatoria semestral"],
        ["1", "Estudiante FIEI", "Tarea de Usuario", "1. Realizar matrícula web y completar ficha socioacadémica (Sem. 0)", "Ficha digital (horas laborales)"],
        ["2", "Plataforma BI+IA", "Tarea de Servicio", "2. Ejecutar pipeline ETL SSIS: ingesta y limpieza al DataMart (Sem. 1)", "DataMart dw (SQL Server)"],
        ["3", "Plataforma BI+IA", "Tarea de Script", "3. Procesar modelo predictivo M1 en Python (Random Forest + K-Means) (Sem. 1-2)", "P(Drop)M1 y Clúster"],
        ["G1", "Plataforma BI+IA", "Compuerta XOR", "¿Riesgo M1 detectado? [P(Drop) >= 0.50]", "Bifurcación Bajo / Alto Riesgo"],
        ["4", "Dir. de Tutoría", "Tarea de Usuario", "4. Consultar Dashboard Power BI y prescribir plan diferenciado (Sem. 2-3)", "Ficha de rescate por clúster"],
        ["5", "Estudiante FIEI", "Tarea de Usuario", "5. Recibir notificación temprana y plan de acompañamiento personalizado (Sem. 3)", "Notificación digital"],
        ["6", "Estudiante FIEI", "Tarea de Usuario", "6. Asistir a clases y rendir evaluaciones continuas (Sem. 4-8)", "Participación activa"],
        ["7", "Docente / Tutor", "Tarea de Usuario", "7. Registrar notas de prácticas en intranet centralizada en vivo (Sem. 4-8)", "Intranet sincronizada"],
        ["8", "Plataforma BI+IA", "Tarea de Servicio", "8. Re-evaluación adaptativa M2: procesar Cubo SSAS y varianza Delta P (Sem. 9-10)", "Cubo OLAP y Delta P"],
        ["G2", "Plataforma BI+IA", "Compuerta XOR", "¿Estado de evolución en M2?", "Recuperado vs Crítico persistente"],
        ["9", "Dir. de Tutoría", "Tarea de Usuario", "9. Ejecutar simulación What-If en Power BI y redistribuir tutores (Sem. 11-14)", "Parámetros What-If"],
        ["10", "Docente / Tutor", "Tarea de Usuario", "10. Impartir asesorías focalizadas y monitorías de rescate (Sem. 11-15)", "Tutoría adaptativa"],
        ["11", "Estudiante FIEI", "Tarea de Usuario", "11. Rendir evaluaciones finales con soporte efectivo (Semana 16)", "Pruebas finales aprobadas"],
        ["12", "Dir. de Tutoría", "Tarea de Usuario", "12. Auditar retención de cohorte y generar reporte SINEACE/SUNEDU (Sem. 16)", "Informe ejecutivo de auditoría"],
        ["Fin", "Dir. de Tutoría", "Fin (Círculo rojo)", "Cohorte culminada: Retención maximizada y gratuidad preservada", "Impacto institucional"]
    ]
    add_table_apa(doc, "3", "Inventario Detallado de Elementos del Proceso Optimizado (TO-BE)", to_be_headers, to_be_rows, [0.4, 1.3, 1.3, 2.2, 1.3])
    
    add_h3(doc, "Anotaciones Obligatorias de Valor Agregado (BI + IA)")
    add_p_apa(doc, "El modelador debe vincular las siguientes cinco anotaciones BPMN en recuadros de color verde claro (#DCFCE7) con línea punteada hacia sus respectivas actividades:", first_indent=False)
    
    add_callout_box(doc, "✓ VALOR AGREGADO 1: Detección Preventiva Temprana (Conectar a Tarea 3)", 
                    "Diagnóstico multivariable en Semana 2 al ingreso, calculando la probabilidad de deserción antes de que el estudiante rinda su primer examen parcial.",
                    bg_color="DCFCE7", border_color="16A34A", text_color="166534")
                    
    add_callout_box(doc, "✓ VALOR AGREGADO 2: Prescripción Diferenciada por Clúster (Conectar a Tarea 4)", 
                    "Tutoría personalizada según perfil identificado por K-Means: nivelación en ciencias básicas, flexibilización por jornada laboral o gestión de becas y comedor.",
                    bg_color="DCFCE7", border_color="16A34A", text_color="166534")
                    
    add_callout_box(doc, "✓ VALOR AGREGADO 3: Erradicación de Silos de Datos (Conectar a Tarea 7)", 
                    "Eliminación de archivos Excel en computadoras docentes mediante digitación continua obligatoria en intranet sincronizada con el DataMart.",
                    bg_color="DCFCE7", border_color="16A34A", text_color="166534")
                    
    add_callout_box(doc, "✓ VALOR AGREGADO 4: Evaluación Adaptativa Dinámica (Conectar a Tarea 8)", 
                    "Recálculo dinámico del riesgo en mitad de semestre (M2) a través de cubos SSAS, cuantificando la varianza Delta P para detectar recuperaciones o alertas críticas.",
                    bg_color="DCFCE7", border_color="16A34A", text_color="166534")
                    
    add_callout_box(doc, "✓ VALOR AGREGADO 5: Simulación Decisional What-If (Conectar a Tarea 9)", 
                    "Ajuste interactivo de umbrales y redistribución de horas docentes en Power BI para maximizar la cobertura tutorial en las semanas decisivas del ciclo.",
                    bg_color="DCFCE7", border_color="16A34A", text_color="166534")

    # CRITERIOS DE REVISIÓN Y ENTREGA
    add_h2(doc, "4. Parámetros de Diseño y Protocolo de Exportación")
    add_p_apa(doc,
        "Para garantizar una apariencia profesional y uniforme en el informe impreso y digital, se deben respetar los siguientes parámetros estéticos dentro de Bizagi Process Modeler:",
        first_indent=False
    )
    
    check_headers = ["Aspecto de Diseño", "Directriz Requerida"]
    check_rows = [
        ["Tipografía", "Segoe UI o Arial, 8 a 9 puntos para tareas, 11 puntos en negrita para encabezados de Pools."],
        ["Conectores de Flujo", "Ortogonales (ángulos rectos a 90°). Mantener alineación horizontal y vertical perfecta."],
        ["Color del Proceso AS-IS", "Pool en tono gris/azul pizarra. Tareas 3, 5, 6 y 7 con relleno tenue rojizo (#FEE2E2) y borde rojo para resaltar cuellos de botella."],
        ["Color del Proceso TO-BE", "Carril 4 (Plataforma BI + IA) con relleno tenue verdoso (#DCFCE7) y borde verde oscuro (#16A34A) para destacar la automatización."],
        ["Exportación de Imagen", "Pestaña 'Publicar' / 'Exportar' -> Imagen -> Formato PNG -> Resolución 300 DPI (Obligatorio para evitar pixelado en Word)."],
        ["Nombres de Archivos", "AS_IS_BPMN_PROYECTO_GADI.png\nTO_BE_BPMN_PROYECTO_GADI.png\nMODELADO_PROCESOS_GADI_BIZAGI.bpm"]
    ]
    add_table_apa(doc, "4", "Checklist de Criterios Técnicos para el Modelador", check_headers, check_rows, [2.2, 4.3])
    
    out_docx = os.path.abspath("GUIA_MODELADO_BIZAGI_AS_IS_Y_TO_BE.docx")
    doc.save(out_docx)
    print(f"DOCX generado con éxito en: {out_docx}")
    return out_docx

if __name__ == "__main__":
    build_guia_docx()
